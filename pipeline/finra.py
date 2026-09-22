from __future__ import annotations

import csv
import io
import json
import re
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from .methodology import period_pct_change, to_observations


SOURCE_URL = "https://www.finra.org/rules-guidance/key-topics/margin-accounts/margin-statistics"
LEGACY_SPLIT_CUTOFF = "2010-01-31"


class FinraError(RuntimeError):
    pass


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def _parse_amount(raw):
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip()
    if not s or s.lower() in {"na", "n/a", "-", "."}:
        return None
    negative = s.startswith("(") and s.endswith(")")
    s = s.replace("$", "").replace(",", "").replace("(", "").replace(")", "")
    value = float(s)
    return -value if negative else value


def _business_month_end(value) -> str:
    if isinstance(value, datetime):
        d = value.date()
    elif isinstance(value, date):
        d = value
    else:
        s = str(value).strip()
        parsers = [
            "%b-%y",
            "%b %Y",
            "%Y-%m",
            "%Y-%m-%d",
            "%m/%d/%Y",
        ]
        d = None
        for fmt in parsers:
            try:
                d = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            raise FinraError(f"Unrecognized month/date value: {value!r}")

    last = date(d.year, d.month, monthrange(d.year, d.month)[1])
    while last.weekday() >= 5:
        last -= timedelta(days=1)
    return last.isoformat()


def _first_matching_header(headers: list[str], predicate):
    for header in headers:
        if predicate(_norm(header)):
            return header
    return None


def _match_columns(headers: list[str]) -> dict[str, str | None]:
    """
    Match both current FINRA columns and the pre-Feb-2010 combined-credit layout.

    FINRA's page states that through January 2010 cash- and margin-account free
    credits were combined. The parser therefore treats the individual components
    as unavailable for that legacy period and preserves a total-free-credit field.
    """
    month = _first_matching_header(
        headers,
        lambda h: h in {"month", "date", "year month", "month year"} or "month year" in h,
    )
    debt = _first_matching_header(
        headers,
        lambda h: (
            "debit" in h
            and "balance" in h
            and ("margin" in h or "customer" in h)
        ) or h == "margin debt",
    )

    combined = _first_matching_header(
        headers,
        lambda h: "free credit" in h and "cash" in h and "margin" in h,
    )
    cash = _first_matching_header(
        headers,
        lambda h: "free credit" in h and "cash" in h and "margin" not in h,
    )
    margin = _first_matching_header(
        headers,
        lambda h: "free credit" in h and "margin" in h and "cash" not in h,
    )

    # Some legacy exports use a generic "Free Credit Balances" / "Credit Balances"
    # header rather than spelling out both account types.
    if combined is None:
        combined = _first_matching_header(
            headers,
            lambda h: (
                "credit balance" in h
                and "debit" not in h
                and "cash" not in h
                and "margin" not in h
            ),
        )

    if month is None or debt is None:
        raise FinraError(f"Missing FINRA month/debit columns; headers={headers}")
    if combined is None and cash is None and margin is None:
        raise FinraError(f"Missing FINRA free-credit columns; headers={headers}")

    return {
        "month": month,
        "margin_debt": debt,
        "cash_free_credit": cash,
        "margin_free_credit": margin,
        "combined_free_credit": combined,
    }


def _normalize_row(raw_date, margin_debt, cash_free_credit, margin_free_credit, combined_free_credit):
    obs_date = _business_month_end(raw_date)
    debt = _parse_amount(margin_debt)
    cash = _parse_amount(cash_free_credit)
    margin = _parse_amount(margin_free_credit)
    combined = _parse_amount(combined_free_credit)

    # FINRA explicitly documents that through Jan-2010 free-credit balances in
    # cash and margin accounts were combined. If an old workbook places that
    # combined value in one of today's component columns, move it into the total
    # rather than silently mislabeling it as a cash or margin component.
    if obs_date <= LEGACY_SPLIT_CUTOFF:
        if combined is None:
            non_null_components = [v for v in (cash, margin) if v is not None]
            if len(non_null_components) == 1:
                combined = non_null_components[0]
                cash = None
                margin = None
            elif len(non_null_components) == 2:
                combined = cash + margin
                cash = None
                margin = None
        else:
            cash = None
            margin = None

    total = combined
    if total is None and cash is not None and margin is not None:
        total = cash + margin

    return {
        "date": obs_date,
        "margin_debt": debt,
        "cash_free_credit": cash,
        "margin_free_credit": margin,
        "combined_free_credit": combined,
        "total_free_credit": total,
    }


def _parse_mapping_row(row: dict, cols: dict[str, str | None]):
    return _normalize_row(
        row.get(cols["month"]),
        row.get(cols["margin_debt"]),
        row.get(cols["cash_free_credit"]) if cols["cash_free_credit"] else None,
        row.get(cols["margin_free_credit"]) if cols["margin_free_credit"] else None,
        row.get(cols["combined_free_credit"]) if cols["combined_free_credit"] else None,
    )


def parse_finra_csv(text: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise FinraError("FINRA CSV has no header")
    cols = _match_columns(reader.fieldnames)

    rows = []
    for row in reader:
        month_col = cols["month"]
        if not month_col or not str(row.get(month_col) or "").strip():
            continue
        rows.append(_parse_mapping_row(row, cols))

    return _deduplicate_rows(rows)


def _cell(values, idx):
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def _row_quality(row: dict) -> int:
    return sum(
        row.get(k) is not None
        for k in (
            "margin_debt",
            "cash_free_credit",
            "margin_free_credit",
            "combined_free_credit",
            "total_free_credit",
        )
    )


def _deduplicate_rows(rows: list[dict]) -> list[dict]:
    by_date: dict[str, dict] = {}
    for row in rows:
        current = by_date.get(row["date"])
        if current is None or _row_quality(row) >= _row_quality(current):
            by_date[row["date"]] = row
    return [by_date[d] for d in sorted(by_date)]


def parse_finra_xlsx(path: str | Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FinraError("openpyxl is required for XLSX import; install requirements.txt") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    all_rows = []

    for ws in wb.worksheets:
        data = list(ws.iter_rows(values_only=True))
        for header_idx, row in enumerate(data[:30]):
            headers = ["" if v is None else str(v) for v in row]
            try:
                cols = _match_columns(headers)
            except FinraError:
                continue

            indexes = {
                key: (headers.index(header) if header is not None else None)
                for key, header in cols.items()
            }

            parsed = []
            for values in data[header_idx + 1 :]:
                month_idx = indexes["month"]
                if month_idx is None or month_idx >= len(values) or values[month_idx] is None:
                    continue
                try:
                    parsed.append(
                        _normalize_row(
                            _cell(values, indexes["month"]),
                            _cell(values, indexes["margin_debt"]),
                            _cell(values, indexes["cash_free_credit"]),
                            _cell(values, indexes["margin_free_credit"]),
                            _cell(values, indexes["combined_free_credit"]),
                        )
                    )
                except (ValueError, FinraError):
                    continue

            if parsed:
                all_rows.extend(parsed)
                break

    rows = _deduplicate_rows(all_rows)
    if not rows:
        raise FinraError("Could not locate FINRA margin-statistics table in XLSX")
    return rows


def _freshness(as_of: str | None, fetched_at: datetime, max_age_days: int = 70):
    if not as_of:
        return "missing", None
    age = (fetched_at.date() - date.fromisoformat(as_of)).days
    return ("fresh" if age <= max_age_days else "stale"), max(age, 0)


def _raw_metric(
    metric_id: str,
    name: str,
    units: str,
    rows: list[dict],
    field: str,
    fetched_at: datetime,
    *,
    polarity: str = "contextual",
    description: str = "FINRA monthly margin-statistics series.",
) -> dict:
    obs = [
        {
            "date": r["date"],
            "value": r[field],
            "status": "observed" if r[field] is not None else "missing",
        }
        for r in rows
    ]
    present = [o for o in obs if o["value"] is not None]
    latest = present[-1] if present else None
    as_of = latest["date"] if latest else None
    state, age = _freshness(as_of, fetched_at)

    return {
        "schema_version": "1.0.0",
        "environment": "production",
        "metric": {
            "id": metric_id,
            "name": name,
            "description": description,
            "pillar": "leverage",
            "units": units,
            "frequency": "monthly",
            "polarity": polarity,
        },
        "source": {
            "provider": "FINRA",
            "dataset": "Margin Statistics",
            "series_id": None,
            "url": SOURCE_URL,
            "license_note": (
                "Official FINRA monthly margin statistics. Through Jan-2010, "
                "cash- and margin-account free credit was reported as a combined amount."
            ),
            "redistribution": "unknown",
        },
        "coverage": {
            "history_start": obs[0]["date"] if obs else None,
            "history_end": obs[-1]["date"] if obs else None,
            "timezone": "America/New_York",
            "expected_observation_lag_days": 25,
        },
        "freshness": {
            "state": state,
            "max_age_days": 70,
            "age_days": age,
            "evaluated_at": fetched_at.isoformat().replace("+00:00", "Z"),
            "reason": None,
        },
        "lineage": {
            "kind": "raw",
            "inputs": [],
            "formula": None,
            "transform_version": "finra-raw-v2",
        },
        "baselines": [
            {
                "id": "expanding-pit",
                "type": "full_history_percentile",
                "window_observations": None,
                "min_observations": 36,
                "point_in_time": True,
                "notes": "Strict-past expanding baseline.",
            },
            {
                "id": "rolling-120m",
                "type": "rolling_percentile",
                "window_observations": 120,
                "min_observations": 36,
                "point_in_time": True,
                "notes": "Trailing ten-year monthly baseline.",
            },
        ],
        "latest": {
            "as_of": as_of,
            "fetched_at": fetched_at.isoformat().replace("+00:00", "Z"),
            "value": latest["value"] if latest else None,
            "revision_tag": None,
        },
        "observations": obs,
    }


def _derived_pct(metric_id: str, name: str, source_metric: dict, periods: int, formula: str) -> dict:
    derived_obs = period_pct_change(source_metric["observations"], periods)
    present = [o for o in derived_obs if o["value"] is not None]
    latest = present[-1] if present else None

    result = json.loads(json.dumps(source_metric))
    result["metric"].update(
        {"id": metric_id, "name": name, "units": "percent", "polarity": "contextual"}
    )
    result["lineage"] = {
        "kind": "derived",
        "inputs": [source_metric["metric"]["id"]],
        "formula": formula,
        "transform_version": "pct-change-v1",
    }
    result["latest"]["as_of"] = latest["date"] if latest else None
    result["latest"]["value"] = latest["value"] if latest else None
    result["observations"] = to_observations(derived_obs)
    return result


def _ratio_metric(rows: list[dict], base_metric: dict) -> dict:
    obs = []
    for row in rows:
        total = row.get("total_free_credit")
        debt = row.get("margin_debt")
        value = None if total in (None, 0) or debt is None else debt / total
        obs.append(
            {
                "date": row["date"],
                "value": value,
                "status": "observed" if value is not None else "missing",
            }
        )

    present = [o for o in obs if o["value"] is not None]
    latest = present[-1] if present else None

    result = json.loads(json.dumps(base_metric))
    result["metric"].update(
        {
            "id": "margin_debt_to_free_credit",
            "name": "Margin Debt / Total Free Credit",
            "units": "ratio",
            "polarity": "higher_is_riskier",
            "description": (
                "Margin debt divided by total free credit. Uses the combined legacy "
                "free-credit series through Jan-2010 and cash+margin thereafter."
            ),
        }
    )
    result["lineage"] = {
        "kind": "derived",
        "inputs": ["finra_margin_debt", "finra_total_free_credit"],
        "formula": "margin_debt / total_free_credit",
        "transform_version": "ratio-v2",
    }
    result["latest"]["as_of"] = latest["date"] if latest else None
    result["latest"]["value"] = latest["value"] if latest else None
    result["observations"] = obs
    return result


def build_finra_metrics(rows: list[dict], fetched_at: datetime | None = None) -> dict[str, dict]:
    fetched_at = fetched_at or datetime.now(timezone.utc)

    debt = _raw_metric(
        "finra_margin_debt",
        "FINRA Margin Debt",
        "USD millions",
        rows,
        "margin_debt",
        fetched_at,
    )
    total = _raw_metric(
        "finra_total_free_credit",
        "FINRA Total Free Credit",
        "USD millions",
        rows,
        "total_free_credit",
        fetched_at,
        description=(
            "Total FINRA free credit. Legacy combined series through Jan-2010; "
            "cash + margin-account free credit thereafter."
        ),
    )
    cash = _raw_metric(
        "finra_cash_free_credit",
        "FINRA Cash Account Free Credit",
        "USD millions",
        rows,
        "cash_free_credit",
        fetched_at,
        description="Cash-account free credit; separate component available from Feb-2010 onward.",
    )
    margin = _raw_metric(
        "finra_margin_free_credit",
        "FINRA Margin Account Free Credit",
        "USD millions",
        rows,
        "margin_free_credit",
        fetched_at,
        description="Margin-account free credit; separate component available from Feb-2010 onward.",
    )
    mom = _derived_pct(
        "finra_margin_debt_mom_pct",
        "FINRA Margin Debt MoM",
        debt,
        1,
        "(x_t / x_t-1 - 1) * 100",
    )
    yoy = _derived_pct(
        "finra_margin_debt_yoy_pct",
        "FINRA Margin Debt YoY",
        debt,
        12,
        "(x_t / x_t-12 - 1) * 100",
    )
    ratio = _ratio_metric(rows, debt)

    return {
        m["metric"]["id"]: m
        for m in [debt, total, cash, margin, mom, yoy, ratio]
    }
